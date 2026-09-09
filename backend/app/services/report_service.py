import io
from datetime import date
import uuid

from sqlalchemy.orm import Session

from app.models.project import Project
from app.models.requirement import Requirement
from app.models.feature import Feature
from app.models.task import Task
from app.models.role import Role
from app.models.risk import Risk
from app.models.estimate import Estimate
from app.models.tech_stack import TechStackRecommendation
from app.repositories import (
    requirement_repository,
    feature_repository,
    risk_repository,
    estimate_repository,
    tech_stack_repository,
)
from app.services.mvp_service import generate_mvp_recommendation


class ReportGenerationError(Exception):
    pass


def gather_report_data(db: Session, project_id: uuid.UUID) -> dict:
    project = db.query(Project).filter(Project.id == project_id).first()
    if not project:
        raise ReportGenerationError("Project not found")

    requirements = db.query(Requirement).filter(
        Requirement.project_id == project_id
    ).all()
    features = feature_repository.list_features(db, project_id)

    feature_ids = [f.id for f in features]
    role_names = {r.id: r.name for r in db.query(Role).all()}

    tasks = []
    if feature_ids:
        tasks_q = (
            db.query(Task)
            .filter(Task.feature_id.in_(feature_ids))
            .all()
        )
        for t in tasks_q:
            tasks.append(
                {
                    "title": t.title,
                    "role": role_names.get(t.role_id, "Unassigned"),
                    "base_hours": t.base_hours,
                }
            )

    estimate = estimate_repository.get_estimate(db, project_id)
    risks = db.query(Risk).filter(Risk.project_id == project_id).all()
    tech = tech_stack_repository.get_recommendation(db, project_id)
    mvp = generate_mvp_recommendation(db, project_id)

    tech_stack = []
    if tech and tech.stack:
        if isinstance(tech.stack, list):
            tech_stack = tech.stack
        elif isinstance(tech.stack, dict):
            tech_stack = [
                {"category": k, "recommendation": v}
                for k, v in tech.stack.items()
            ]

    return {
        "project": project,
        "requirements": requirements,
        "features": features,
        "tasks": tasks,
        "estimate": estimate,
        "risks": risks,
        "tech_stack": tech_stack,
        "mvp": mvp,
    }


def _section_heading_hours(estimate):
    if estimate is None:
        return {
            "min_hours": None,
            "expected_hours": None,
            "max_hours": None,
            "min_cost": None,
            "expected_cost": None,
            "max_cost": None,
            "timeline_weeks": None,
        }
    return {
        "min_hours": estimate.min_hours,
        "expected_hours": estimate.expected_hours,
        "max_hours": estimate.max_hours,
        "min_cost": estimate.min_cost,
        "expected_cost": estimate.expected_cost,
        "max_cost": estimate.max_cost,
        "timeline_weeks": estimate.timeline_weeks,
    }


# ------------------------------------------------------------------
# PDF
# ------------------------------------------------------------------

def _wrap(text, size=9, max_chars=95):
    text = str(text)
    if len(text) <= max_chars:
        return [text]
    words = text.split()
    lines = []
    current = ""
    for w in words:
        candidate = (current + " " + w).strip()
        if len(candidate) <= max_chars:
            current = candidate
        else:
            if current:
                lines.append(current)
            current = w
    if current:
        lines.append(current)
    return lines


def generate_pdf(data: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.units import mm
    from reportlab.lib import colors
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        HRFlowable,
        ListFlowable,
        ListItem,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

    project = data["project"]
    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "TitleCustom", parent=styles["Title"], fontSize=24, spaceAfter=6
    )
    h2 = ParagraphStyle(
        "H2Custom", parent=styles["Heading2"], fontSize=14, spaceBefore=10, spaceAfter=4
    )
    body = ParagraphStyle(
        "BodyCustom", parent=styles["BodyText"], fontSize=9.5, leading=13
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=18 * mm,
        rightMargin=18 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"Project Plan - {project.title}",
    )

    story = []
    story.append(Paragraph("ProjectScope AI", title_style))
    story.append(Paragraph("Project Plan", h2))
    story.append(
        Paragraph(
            f"{project.title} &middot; Generated {date.today().isoformat()}",
            body,
        )
    )
    story.append(HRFlowable(width="100%", thickness=1, color=colors.grey))
    story.append(Spacer(1, 6))
    story.append(
        Paragraph(
            f"<b>Status:</b> {project.status} &nbsp;&nbsp;<b>Platform:</b> "
            f"{project.platform or 'not specified'} &nbsp;&nbsp;<b>Budget:</b> "
            f"{project.budget or 'not specified'}",
            body,
        )
    )
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"<b>Description:</b> {project.description}", body))

    # Executive summary
    story.append(Spacer(1, 6))
    story.append(Paragraph("Executive Summary", h2))
    est = _section_heading_hours(data["estimate"])
    if est["expected_hours"]:
        exec_lines = [
            f"This project ({project.title}) has {len(data['features'])} extracted "
            f"features and is estimated at an expected {est['expected_hours']:.0f} "
            f"hours of effort (range {est['min_hours']:.0f}-{est['max_hours']:.0f} hours).",
            f"Expected cost is ${est['expected_cost']:,.0f} "
            f"(range ${est['min_cost']:,.0f}-${est['max_cost']:,.0f}).",
            f"Estimated delivery timeline is approximately "
            f"{est['timeline_weeks']:.1f} weeks.",
        ]
    else:
        exec_lines = [
            "This project has not been fully estimated yet. Run the analysis and "
            "estimation pipeline to see effort, cost and timeline figures."
        ]
    for line in exec_lines:
        story.append(Paragraph(line, body))

    # Assumptions / limitations
    story.append(Spacer(1, 6))
    story.append(Paragraph("Assumptions &amp; Limitations", h2))
    story.append(
        Paragraph(
            "These estimates are generated from the project description, extracted "
            "features and deterministic estimation rules combined with an ML model. "
            "They are planning aids, not guarantees. Figures assume a standard "
            "engineering team and do not include unforeseen scope creep.",
            body,
        )
    )

    # Requirements
    story.append(Spacer(1, 6))
    story.append(Paragraph("Requirements", h2))
    if data["requirements"]:
        req_table = [
            ["#", "Category", "Requirement"]
        ]
        for i, r in enumerate(data["requirements"], start=1):
            req_table.append([i, r.category, r.description])
        t = Table(req_table, colWidths=[20, 70, 150])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(t)
    else:
        story.append(Paragraph("No requirements extracted.", body))

    # Features
    story.append(Spacer(1, 6))
    story.append(Paragraph("Features", h2))
    if data["features"]:
        feat_rows = [["Feature", "Priority", "Complexity"]]
        for f in data["features"]:
            feat_rows.append(
                [
                    f.canonical_name.lower().replace("_", " "),
                    f.priority,
                    f.complexity,
                ]
            )
        t = Table(feat_rows, colWidths=[140, 50, 50])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(t)
    else:
        story.append(Paragraph("No features extracted.", body))

    # MVP
    story.append(Spacer(1, 6))
    story.append(Paragraph("Recommended MVP", h2))
    mvp = data["mvp"]
    mvp_features = [f["canonical_name"] for f in mvp.get("mvp_features", [])]
    story.append(
        Paragraph(
            "Core MVP features: "
            + (", ".join(f.lower().replace("_", " ") for f in mvp_features)
               if mvp_features else "none identified."),
            body,
        )
    )
    story.append(Paragraph(mvp.get("reasoning", ""), body))

    # Team / roles + tasks
    story.append(Spacer(1, 6))
    story.append(Paragraph("Team &amp; Task Breakdown", h2))
    if data["tasks"]:
        task_rows = [["Task", "Role", "Est. Hours"]]
        for t in data["tasks"]:
            task_rows.append([t["title"], t["role"], t["base_hours"]])
        t = Table(task_rows, colWidths=[150, 70, 50])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(t)
    else:
        story.append(Paragraph("No tasks generated yet.", body))

    # Estimate
    story.append(Spacer(1, 6))
    story.append(Paragraph("Effort, Cost &amp; Timeline", h2))
    if est["expected_hours"]:
        est_rows = [
            ["Metric", "Min", "Expected", "Max"],
            [
                "Effort (hours)",
                f"{est['min_hours']:.0f}",
                f"{est['expected_hours']:.0f}",
                f"{est['max_hours']:.0f}",
            ],
            [
                "Cost",
                f"${est['min_cost']:,.0f}",
                f"${est['expected_cost']:,.0f}",
                f"${est['max_cost']:,.0f}",
            ],
            ["Timeline (weeks)", "-", f"{est['timeline_weeks']:.1f}", "-"],
        ]
        t = Table(est_rows, colWidths=[90, 60, 60, 60])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(t)
    else:
        story.append(Paragraph("Not estimated yet. Run estimation first.", body))

    # Risks
    story.append(Spacer(1, 6))
    story.append(Paragraph("Risks", h2))
    if data["risks"]:
        risk_rows = [
            ["Risk", "Probability", "Impact", "Mitigation"]
        ]
        for r in data["risks"]:
            risk_rows.append(
                [
                    r.description,
                    r.probability,
                    r.impact,
                    r.mitigation or "",
                ]
            )
        t = Table(risk_rows, colWidths=[90, 50, 40, 100])
        t.setStyle(
            TableStyle(
                [
                    ("GRID", (0, 0), (-1, -1), 0.3, colors.grey),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                ]
            )
        )
        story.append(t)
    else:
        story.append(Paragraph("No risks identified.", body))

    # Technology
    story.append(Spacer(1, 6))
    story.append(Paragraph("Technology Stack", h2))
    if data["tech_stack"]:
        stack_items = []
        for item in data["tech_stack"]:
            if isinstance(item, dict):
                cat = item.get("category", "stack")
                rec = item.get("recommendation", item.get("reason", ""))
                stack_items.append(f"<b>{cat}:</b> {rec}")
        if stack_items:
            story.append(ListFlowable([
                ListItem(Paragraph(s, body)) for s in stack_items
            ]))
        else:
            story.append(Paragraph("No technology recommendation recorded.", body))
    else:
        story.append(Paragraph("No technology recommendation recorded.", body))

    doc.build(story)
    return buf.getvalue()


# ------------------------------------------------------------------
# DOCX
# ------------------------------------------------------------------

def generate_docx(data: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    project = data["project"]
    document = Document()

    styles = document.styles
    normal = styles["Normal"]
    normal.font.size = Pt(10.5)

    title = document.add_heading(project.title, level=0)
    sub = document.add_paragraph(
        f"ProjectScope AI - Project Plan - Generated {date.today().isoformat()}"
    )
    sub.alignment = WD_ALIGN_PARAGRAPH.LEFT

    p = document.add_paragraph()
    p.add_run("Status: ").bold = True
    p.add_run(project.status)
    p.add_run("    Platform: ").bold = True
    p.add_run(project.platform or "not specified")
    p.add_run("    Budget: ").bold = True
    p.add_run(project.budget or "not specified")

    document.add_paragraph()
    doc_p = document.add_paragraph()
    doc_p.add_run("Description: ").bold = True
    doc_p.add_run(project.description)

    est = _section_heading_hours(data["estimate"])

    document.add_heading("Executive Summary", level=1)
    if est["expected_hours"]:
        document.add_paragraph(
            f"This project has {len(data['features'])} extracted features and is "
            f"estimated at an expected {est['expected_hours']:.0f} hours "
            f"(range {est['min_hours']:.0f}-{est['max_hours']:.0f} hours)."
        )
        document.add_paragraph(
            f"Expected cost is ${est['expected_cost']:,.0f} "
            f"(range ${est['min_cost']:,.0f}-${est['max_cost']:,.0f}), with an "
            f"estimated timeline of approximately {est['timeline_weeks']:.1f} weeks."
        )
    else:
        document.add_paragraph(
            "This project has not been fully estimated yet. Run the analysis and "
            "estimation pipeline first."
        )

    document.add_heading("Assumptions & Limitations", level=1)
    document.add_paragraph(
        "These estimates are generated from the project description, extracted "
        "features and deterministic estimation rules combined with an ML model. "
        "They are planning aids, not guarantees."
    )

    document.add_heading("Requirements", level=1)
    if data["requirements"]:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "#"
        hdr[1].text = "Category"
        hdr[2].text = "Requirement"
        for i, r in enumerate(data["requirements"], start=1):
            cells = table.add_row().cells
            cells[0].text = str(i)
            cells[1].text = r.category
            cells[2].text = r.description
    else:
        document.add_paragraph("No requirements extracted.")

    document.add_heading("Features", level=1)
    if data["features"]:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Feature"
        hdr[1].text = "Priority"
        hdr[2].text = "Complexity"
        for f in data["features"]:
            cells = table.add_row().cells
            cells[0].text = f.canonical_name.lower().replace("_", " ")
            cells[1].text = f.priority
            cells[2].text = f.complexity
    else:
        document.add_paragraph("No features extracted.")

    document.add_heading("Recommended MVP", level=1)
    mvp = data["mvp"]
    mvp_features = [f["canonical_name"] for f in mvp.get("mvp_features", [])]
    document.add_paragraph(
        "Core MVP features: "
        + (", ".join(f.lower().replace("_", " ") for f in mvp_features)
           if mvp_features else "none identified.")
    )
    document.add_paragraph(mvp.get("reasoning", ""))

    document.add_heading("Team & Task Breakdown", level=1)
    if data["tasks"]:
        table = document.add_table(rows=1, cols=3)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Task"
        hdr[1].text = "Role"
        hdr[2].text = "Est. Hours"
        for t in data["tasks"]:
            cells = table.add_row().cells
            cells[0].text = t["title"]
            cells[1].text = t["role"]
            cells[2].text = str(t["base_hours"])
    else:
        document.add_paragraph("No tasks generated yet.")

    document.add_heading("Effort, Cost & Timeline", level=1)
    if est["expected_hours"]:
        table = document.add_table(rows=4, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Min"
        hdr[2].text = "Expected"
        hdr[3].text = "Max"
        rows_vals = [
            ("Effort (hours)", est["min_hours"], est["expected_hours"], est["max_hours"]),
            ("Cost", f"${est['min_cost']:,.0f}", f"${est['expected_cost']:,.0f}", f"${est['max_cost']:,.0f}"),
            ("Timeline (weeks)", "-", f"{est['timeline_weeks']:.1f}", "-"),
        ]
        for i, (a, b, c, d) in enumerate(rows_vals, start=1):
            cells = table.rows[i].cells
            cells[0].text = a
            cells[1].text = f"{b:.0f}" if isinstance(b, float) else str(b)
            cells[2].text = f"{c:.0f}" if isinstance(c, float) else str(c)
            cells[3].text = f"{d:.0f}" if isinstance(d, float) else str(d)
    else:
        document.add_paragraph("Not estimated yet. Run estimation first.")

    document.add_heading("Risks", level=1)
    if data["risks"]:
        table = document.add_table(rows=1, cols=4)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Risk"
        hdr[1].text = "Probability"
        hdr[2].text = "Impact"
        hdr[3].text = "Mitigation"
        for r in data["risks"]:
            cells = table.add_row().cells
            cells[0].text = r.description
            cells[1].text = r.probability
            cells[2].text = r.impact
            cells[3].text = r.mitigation or ""
    else:
        document.add_paragraph("No risks identified.")

    document.add_heading("Technology Stack", level=1)
    if data["tech_stack"]:
        for item in data["tech_stack"]:
            if isinstance(item, dict):
                cat = item.get("category", "stack")
                rec = item.get("recommendation", item.get("reason", ""))
                p = document.add_paragraph()
                p.add_run(f"{cat}: ").bold = True
                p.add_run(str(rec))
    else:
        document.add_paragraph("No technology recommendation recorded.")

    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()
